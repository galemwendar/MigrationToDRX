"""
Генерация документа «Описание структуры промежуточной БД» из dbdiagram.io
"""
import re
from docx import Document
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ─────────────────────────── Парсинг dbdiagram ───────────────────────────────

def parse_dbdiagram(path):
    with open(path, encoding='utf-8') as f:
        content = f.read()

    # Убираем блочные комментарии /* ... */
    content = re.sub(r'/\*.*?\*/', '', content, flags=re.DOTALL)

    tables = []
    # Захватываем каждый Table блок
    for m in re.finditer(
        r'//\s*(.*?)\n\s*Table\s+(\w+)\s*(?:\[note:\s*"([^"]+)"\])?\s*\{([^}]+)\}',
        content, re.DOTALL
    ):
        comment   = m.group(1).strip()
        tbl_name  = m.group(2).strip()
        tbl_note  = m.group(3) or comment
        body      = m.group(4)

        fields = []
        for line in body.splitlines():
            line = line.strip()
            if not line or line.startswith('//'):
                continue
            # Разбираем: name type [attrs]
            fm = re.match(r'(\w+)\s+([\w()]+)(.*)', line)
            if not fm:
                continue
            fname = fm.group(1)
            ftype = fm.group(2)
            attrs = fm.group(3)

            required  = 'not null' in attrs.lower()
            pk        = 'pk' in attrs.lower()
            unique    = 'unique' in attrs.lower()
            note_m    = re.search(r'note:\s*"([^"]+)"', attrs)
            note      = note_m.group(1) if note_m else ''

            fields.append({
                'name':     fname,
                'type':     ftype,
                'required': required,
                'pk':       pk,
                'unique':   unique,
                'note':     note,
            })

        tables.append({
            'name':    tbl_name,
            'title':   tbl_note,
            'comment': comment,
            'fields':  fields,
        })

    return tables


# ─────────────────────────── Хелперы форматирования ──────────────────────────

def set_cell_bg(cell, hex_color):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  hex_color)
    tcPr.append(shd)

def set_cell_borders(cell):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for side in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement(f'w:{side}')
        el.set(qn('w:val'),   'single')
        el.set(qn('w:sz'),    '4')
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), 'BFBFBF')
        tcBorders.append(el)
    tcPr.append(tcBorders)

def bold_run(para, text, size=None, color=None):
    run = para.add_run(text)
    run.bold = True
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*bytes.fromhex(color))
    return run

def normal_run(para, text, size=9, italic=False):
    run = para.add_run(text)
    run.font.size = Pt(size)
    run.italic = italic
    return run


# ─────────────────────────── Генерация DOCX ──────────────────────────────────

HEADER_COLOR = '1F4E79'   # тёмно-синий (заголовок таблицы)
HEADER_TEXT  = 'FFFFFF'
ROW_ALT      = 'EBF3FB'   # чередующаяся строка
ROW_NORM     = 'FFFFFF'
PK_COLOR     = 'FFF2CC'   # жёлтый для PK

COL_WIDTHS = [Cm(4.5), Cm(3.2), Cm(1.8), Cm(1.8), Cm(6.7)]  # Name|Type|Req|PK|Note

def add_table_section(doc, table_info, idx):
    # ── Заголовок секции ──────────────────────────────────────────────────────
    h = doc.add_heading(level=2)
    h.clear()
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after  = Pt(4)
    bold_run(h, f'{idx}. {table_info["name"]}', size=12, color='1F4E79')

    # Подпись (русское название)
    if table_info['title']:
        sub = doc.add_paragraph()
        sub.paragraph_format.space_before = Pt(0)
        sub.paragraph_format.space_after  = Pt(6)
        normal_run(sub, table_info['title'], size=9, italic=True)

    # ── Таблица полей ────────────────────────────────────────────────────────
    tbl = doc.add_table(rows=1, cols=5)
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl.style = 'Table Grid'

    # Ширины столбцов
    for ci, w in enumerate(COL_WIDTHS):
        for row in tbl.rows:
            row.cells[ci].width = w

    # Шапка
    hdr_cells = tbl.rows[0].cells
    headers = ['Поле', 'Тип', 'NOT NULL', 'PK', 'Описание']
    for ci, h_text in enumerate(headers):
        set_cell_bg(hdr_cells[ci], HEADER_COLOR)
        set_cell_borders(hdr_cells[ci])
        p = hdr_cells[ci].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h_text)
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # Строки полей
    for ri, field in enumerate(table_info['fields']):
        row = tbl.add_row()
        cells = row.cells

        bg = PK_COLOR if field['pk'] else (ROW_ALT if ri % 2 else ROW_NORM)
        for ci in range(5):
            set_cell_bg(cells[ci], bg.replace('#', ''))
            set_cell_borders(cells[ci])
            cells[ci].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

        # Имя поля
        p0 = cells[0].paragraphs[0]
        r0 = p0.add_run(field['name'])
        r0.font.size = Pt(9)
        r0.bold = field['pk']
        if field['pk']:
            r0.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

        # Тип
        p1 = cells[1].paragraphs[0]
        r1 = p1.add_run(field['type'])
        r1.font.size = Pt(8.5)
        r1.font.color.rgb = RGBColor(0x26, 0x5E, 0x0B)  # зелёный для типа

        # NOT NULL
        cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = cells[2].paragraphs[0].add_run('✓' if field['required'] else '')
        r2.font.size = Pt(9)
        if field['required']:
            r2.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)

        # PK
        cells[3].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        r3 = cells[3].paragraphs[0].add_run('✓' if field['pk'] else '')
        r3.font.size = Pt(9)
        if field['pk']:
            r3.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

        # Описание
        p4 = cells[4].paragraphs[0]
        desc_parts = []
        if field['note']:
            desc_parts.append(field['note'])
        if field['unique'] and not field['pk']:
            desc_parts.append('UNIQUE')
        r4 = p4.add_run('; '.join(desc_parts))
        r4.font.size = Pt(8.5)

    doc.add_paragraph()  # отступ после таблицы


def generate(diag_path, out_path):
    tables = parse_dbdiagram(diag_path)

    doc = Document()

    # ── Поля страницы ────────────────────────────────────────────────────────
    section = doc.sections[0]
    section.page_width  = Cm(29.7)
    section.page_height = Cm(21.0)   # альбомная
    section.left_margin   = Cm(2.0)
    section.right_margin  = Cm(1.5)
    section.top_margin    = Cm(2.0)
    section.bottom_margin = Cm(1.5)

    # ── Заголовок документа ──────────────────────────────────────────────────
    title_p = doc.add_heading(level=1)
    title_p.clear()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    bold_run(title_p, 'Описание структуры промежуточной базы данных', size=16, color='1F4E79')

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    normal_run(subtitle, 'RusalIntermediate · MS SQL Server · Система миграции документов из PayDox в DRX', size=10, italic=True)

    doc.add_paragraph()

    # ── Вводная таблица: список таблиц ───────────────────────────────────────
    intro_h = doc.add_heading(level=2)
    intro_h.clear()
    bold_run(intro_h, 'Состав базы данных', size=12, color='1F4E79')

    intro_tbl = doc.add_table(rows=1, cols=3)
    intro_tbl.style = 'Table Grid'
    for ci, txt in enumerate(['№', 'Таблица', 'Назначение']):
        c = intro_tbl.rows[0].cells[ci]
        set_cell_bg(c, HEADER_COLOR)
        set_cell_borders(c)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(txt)
        r.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for i, t in enumerate(tables):
        row = intro_tbl.add_row()
        bg = ROW_ALT if i % 2 else ROW_NORM
        for ci in range(3):
            set_cell_bg(row.cells[ci], bg)
            set_cell_borders(row.cells[ci])

        row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        r0 = row.cells[0].paragraphs[0].add_run(str(i + 1))
        r0.font.size = Pt(9)

        r1 = row.cells[1].paragraphs[0].add_run(t['name'])
        r1.font.size = Pt(9)
        r1.bold = True

        r2 = row.cells[2].paragraphs[0].add_run(t['title'])
        r2.font.size = Pt(9)

    doc.add_page_break()

    # ── Детальное описание каждой таблицы ────────────────────────────────────
    detail_h = doc.add_heading(level=1)
    detail_h.clear()
    bold_run(detail_h, 'Детальное описание таблиц', size=14, color='1F4E79')

    for i, t in enumerate(tables):
        add_table_section(doc, t, i + 1)

    doc.save(out_path)
    print(f'Документ сохранён: {out_path}')


if __name__ == '__main__':
    generate(
        r'D:\Applications\Rusal\MigrationToDRX\docs\dbdiagram.io',
        r'D:\Applications\Rusal\MigrationToDRX\docs\DB_Structure.docx',
    )
